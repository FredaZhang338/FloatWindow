from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

PRIMARY = RGBColor(0x2F, 0x54, 0xEB)  # #2F54EB
BG_LIGHT = RGBColor(0xF7, 0xF9, 0xFC)  # #F7F9FC
TEXT_DARK = RGBColor(0x11, 0x11, 0x11)  # #111111
TEXT_BODY = RGBColor(0x59, 0x59, 0x59)  # #595959
ACCENT_GREEN = RGBColor(0x52, 0xC4, 0x1A)  # #52C41A
ACCENT_AMBER = RGBColor(0xFA, 0xAD, 0x14)  # #FAAD14
ACCENT_RED = RGBColor(0xF5, 0x22, 0x2D)   # #F5222D


def set_text(run, text, size, bold=False, color=TEXT_BODY):
    run.text = text
    font = run.font
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color


def add_title(slide, text):
    left = Inches(0.6)
    top = Inches(0.3)
    width = Inches(12.0)
    height = Inches(1.0)
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    set_text(r, text, 38, True, TEXT_DARK)
    
    # subtle subtitle
    p2 = tf.add_paragraph()
    p2.space_before = Pt(6)
    r2 = p2.add_run()
    set_text(r2, "标签体系 × 接口分层 × 同源扩散（立体化视图）", 18, False, TEXT_BODY)


def add_layered_card(slide, left, top, width, height, title, bullets, fill_rgb, line_rgb=None, z=0):
    # Draw a rounded rectangle card with text
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill = shp.fill
    fill.solid()
    fill.fore_color.rgb = fill_rgb
    line = shp.line
    line.color.rgb = line_rgb or fill_rgb
    line.width = Pt(1.5)

    # inner text
    margin = Inches(0.35)
    tbox = slide.shapes.add_textbox(left + margin, top + margin, width - margin*2, height - margin*2)
    tf = tbox.text_frame
    tf.clear()

    # title
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    set_text(r, title, 24, True, PRIMARY)

    for b in bullets:
        p_b = tf.add_paragraph()
        p_b.level = 1
        p_b.space_before = Pt(3)
        r_b = p_b.add_run()
        set_text(r_b, b, 16, False, TEXT_BODY)

    return shp, tbox


def add_kpi(slide, left, top, width, height, label, value, accent=PRIMARY):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    card.line.color.rgb = accent
    card.line.width = Pt(1.5)

    # value
    tbox = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.12), width - Inches(0.4), height - Inches(0.24))
    tf = tbox.text_frame
    tf.clear()
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.LEFT
    r1 = p1.add_run()
    set_text(r1, value, 24, True, TEXT_DARK)

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.LEFT
    r2 = p2.add_run()
    set_text(r2, label, 13, False, TEXT_BODY)

    return card


def build_slide(output_path: str):
    prs = Presentation()
    # Set 16:9 size (13.33" x 7.5")
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    add_title(slide, "账户动账风险｜资金安全用例现状（第8页改版）")

    # Layered cards for 立体化，背->中->前
    total_left = Inches(0.9)
    total_top = Inches(1.4)
    total_width = Inches(11.5)
    total_height = Inches(4.7)

    # back layer (方法论)
    l_back = total_left + Inches(0.35)
    t_back = total_top + Inches(0.35)
    w = total_width
    h = total_height
    back_card, _ = add_layered_card(
        slide,
        l_back,
        t_back,
        w,
        h,
        "方法论 — 不变的建设基座",
        [
            "标签体系：安全风险标签｜军规标签｜用例标签",
            "接口分层：资金流动｜信息流动｜非对账",
            "用例标准：场景化断言，正/反向成对构造，关键断言"
        ],
        RGBColor(0xEB, 0xF0, 0xFF),
        RGBColor(0xE1, 0xE7, 0xFF)
    )

    # middle layer (覆盖维度)
    l_mid = total_left + Inches(0.18)
    t_mid = total_top + Inches(0.18)
    mid_card, _ = add_layered_card(
        slide,
        l_mid,
        t_mid,
        w,
        h,
        "覆盖维度 — 构成立体防线",
        [
            "防错：要素完整性、入参合规、计算口径",
            "防漏：配置/分支覆盖、状态迁移",
            "防重：幂等、重复与补偿处理",
            "防误：异常/时效、跨模块一致性"
        ],
        RGBColor(0xDC, 0xE6, 0xFF),
        RGBColor(0xD0, 0xDC, 0xFF)
    )

    # front layer (结果)
    l_front = total_left
    t_front = total_top
    front_card, _ = add_layered_card(
        slide,
        l_front,
        t_front,
        w,
        h,
        "结果与试点 — 订单交易场景",
        [
            "接口打标示例：fx_exchange_server/fx_exchange_done_service",
            "同源扩散：环境相关一致校验与关键断言",
        ],
        RGBColor(0xFF, 0xFF, 0xFF),
        PRIMARY
    )

    # KPIs on the front card bottom
    kpi_top = t_front + h - Inches(1.25)
    kpi_left = l_front + Inches(0.5)
    kpi_width = Inches(2.6)
    kpi_height = Inches(1.0)
    gap = Inches(0.35)

    add_kpi(slide, kpi_left, kpi_top, kpi_width, kpi_height, "订单交易", "161", accent=PRIMARY)
    add_kpi(slide, kpi_left + (kpi_width + gap), kpi_top, kpi_width, kpi_height, "资金接口识别率", "87.76%", accent=ACCENT_GREEN)
    add_kpi(slide, kpi_left + (kpi_width + gap)*2, kpi_top, kpi_width, kpi_height, "资金接口PO自动化覆盖率", "75.31%", accent=ACCENT_AMBER)
    add_kpi(slide, kpi_left + (kpi_width + gap)*3, kpi_top, kpi_width, kpi_height, "流量转用例", "7", accent=ACCENT_RED)

    prs.save(output_path)


if __name__ == "__main__":
    build_slide("/workspace/slide8_redesign.pptx")
